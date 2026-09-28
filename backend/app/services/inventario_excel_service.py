import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.models.inventario import Inventario, SituacaoDivergencia
from app.models.empresa import Empresa


class InventarioExcelService:
    @staticmethod
    def gerar_excel_inventario(inventario_id: int, user_nome: str = "Administrador") -> io.BytesIO:
        inventario = Inventario.query.get(inventario_id)
        if not inventario:
            raise ValueError("Inventário não encontrado para exportação Excel.")

        empresa = Empresa.query.first()
        nome_empresa = empresa.nome if empresa and empresa.nome else "SABOR IMBATÍVEL"
        nif_empresa = empresa.nif if empresa and empresa.nif else "500000000"

        wb = openpyxl.Workbook()
        wb.remove(wb.active) # Remove a folha inicial por defeito

        # Paleta de Cores & Estilos SIGI ERP
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark Slate
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        section_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
        section_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
        subtitle_font = Font(name="Calibri", size=10, italic=True, color="475569")
        bold_font = Font(name="Calibri", size=11, bold=True)
        regular_font = Font(name="Calibri", size=11)
        
        falta_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Soft Red
        falta_font = Font(name="Calibri", size=11, color="991B1B", bold=True)
        sobra_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Soft Green
        sobra_font = Font(name="Calibri", size=11, color="166534", bold=True)

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        resumo = inventario.calcular_resumo()

        # ==========================================
        # FOLHA 1 — RESUMO EXECUÇÃO & KPIS
        # ==========================================
        ws_resumo = wb.create_sheet(title="Resumo & KPIs")
        ws_resumo.views.sheetView[0].showGridLines = True

        ws_resumo["B2"] = f"{nome_empresa} — SIGI ERP (NIF: {nif_empresa})"
        ws_resumo["B2"].font = title_font
        ws_resumo["B3"] = f"RELATÓRIO DE INVENTÁRIO E RECONCILIAÇÃO FÍSICO-FINANCEIRA — {inventario.numero}"
        ws_resumo["B3"].font = Font(name="Calibri", size=12, bold=True, color="2563EB")
        ws_resumo["B4"] = f"Exportado em: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')} por {user_nome}"
        ws_resumo["B4"].font = subtitle_font

        # Metadados do Inventário
        inv_numero = str(inventario.numero or f"INV-#{inventario.id}")
        armazem_nome = str(inventario.armazem.nome) if inventario.armazem else "Armazém Geral"
        data_inv = inventario.data_inventario.strftime('%d/%m/%Y') if inventario.data_inventario else "-"
        tipo_str = inventario.tipo.value if hasattr(inventario.tipo, 'value') else str(inventario.tipo or "COMPLETO")
        estado_str = inventario.estado.value if hasattr(inventario.estado, 'value') else str(inventario.estado or "APLICADO")
        responsavel_nome = str(inventario.responsavel.name) if inventario.responsavel else "Administrador"

        dados_cabecalho = [
            ("Nº Inventário:", inv_numero),
            ("Armazém / Localização:", armazem_nome),
            ("Data de Referência:", data_inv),
            ("Tipo de Inventário:", tipo_str),
            ("Estado do Processo:", estado_str),
            ("Responsável Técnico:", responsavel_nome),
            ("Observações / Notas:", str(inventario.observacao or "-")),
        ]

        row_idx = 6
        for label, val in dados_cabecalho:
            ws_resumo[f"B{row_idx}"] = label
            ws_resumo[f"B{row_idx}"].font = bold_font
            ws_resumo[f"C{row_idx}"] = val
            ws_resumo[f"C{row_idx}"].font = regular_font
            row_idx += 1

        # Tabela de Indicadores
        row_idx += 1
        ws_resumo[f"B{row_idx}"] = "INDICADORES DE COBERTURA FÍSICA E DESEMPENHO"
        ws_resumo[f"B{row_idx}"].font = Font(name="Calibri", size=12, bold=True, color="1E293B")
        row_idx += 1

        v_sis = float(resumo.get("valor_total_sistema", 0.0))
        v_cnt = float(resumo.get("valor_total_contado", 0.0))
        v_falta = float(resumo.get("valor_estimado_falta", 0.0))
        v_sobra = float(resumo.get("valor_estimado_sobra", 0.0))
        v_liquido = float(resumo.get("ajuste_liquido_financeiro", 0.0))
        acuracia = float(resumo.get("acuracia_percentual", 0.0))

        indicadores = [
            ("Total de Artigos Elegíveis", resumo["total_itens"], "int"),
            ("Artigos Contados", resumo["itens_contados"], "int"),
            ("Artigos Não Contados / Pendentes", resumo["itens_nao_contados"], "int"),
            ("Artigos com Stock Exato (Sem Divergência)", resumo["itens_sem_divergencia"], "int"),
            ("Artigos com Falta de Stock", resumo["itens_com_falta"], "int"),
            ("Artigos com Sobra de Stock", resumo["itens_com_sobra"], "int"),
            ("Quantidade Acumulada em Falta", resumo["quantidade_falta"], "float"),
            ("Quantidade Acumulada em Sobra", resumo["quantidade_sobra"], "float"),
            ("Valor Total do Stock do Sistema (STN)", v_sis, "money"),
            ("Valor Total do Stock Real Contado (STN)", v_cnt, "money"),
            ("Valor Estimado de Faltas (Perda / Quebra) (STN)", v_falta, "money"),
            ("Valor Estimado de Sobras (Excesso) (STN)", v_sobra, "money"),
            ("Ajuste Financeiro Líquido (STN)", v_liquido, "money"),
            ("Índice de Acurácia de Stock (%)", f"{acuracia:.1f}%", "str"),
        ]

        for ind_label, ind_val, fmt in indicadores:
            ws_resumo[f"B{row_idx}"] = ind_label
            ws_resumo[f"B{row_idx}"].font = regular_font
            ws_resumo[f"B{row_idx}"].border = thin_border
            
            c_cell = ws_resumo[f"C{row_idx}"]
            c_cell.font = bold_font
            c_cell.border = thin_border
            c_cell.alignment = Alignment(horizontal="right")

            if fmt == "money":
                c_cell.value = ind_val
                c_cell.number_format = '#,##0.00 "STN"'
            elif fmt == "float":
                c_cell.value = ind_val
                c_cell.number_format = '#,##0.00'
            elif fmt == "int":
                c_cell.value = ind_val
                c_cell.number_format = '#,##0'
            else:
                c_cell.value = str(ind_val)

            row_idx += 1

        # ==========================================
        # FUNÇÃO AUXILIAR PARA PREENCHER TABELAS DE ITENS
        # ==========================================
        def preencher_aba_itens(ws, lista_itens, titulo_aba):
            ws.views.sheetView[0].showGridLines = True
            
            # Cabeçalho da Folha
            ws["A1"] = f"{nome_empresa} — {titulo_aba} (Inventário {inv_numero})"
            ws["A1"].font = title_font
            
            headers = [
                "Código", "Descrição do Artigo", "Tipo", "Unidade",
                "Preço UN (STN)", "Qtd Sistema", "Valor Sistema (STN)",
                "Qtd Contada", "Valor Contado (STN)", "Diferença Qtd",
                "Diferença Valor (STN)", "Situação", "Motivo do Ajuste",
                "Contado Por", "Data/Hora Contagem", "Observação"
            ]
            
            for col_num, h_text in enumerate(headers, 1):
                cell = ws.cell(row=3, column=col_num, value=h_text)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center" if "Qtd" not in h_text and "Valor" not in h_text and "Preço" not in h_text else "right")
                cell.border = thin_border

            curr_row = 4
            for it in lista_itens:
                cod = str(it.obter_codigo() or "-")
                nome = str(it.obter_nome() or "Artigo Desconhecido")
                tipo_item = str(it.obter_tipo_item() or "Artigo")
                un_sigla = str(it.obter_unidade_sigla() or "UN")
                preco_un = float(it.obter_preco_unitario() or 0.0)
                qtd_sis = float(it.quantidade_sistema or 0.0)
                val_sis = round(qtd_sis * preco_un, 2)

                if it.quantidade_contada is not None:
                    qtd_cnt = float(it.quantidade_contada)
                    val_cnt = round(qtd_cnt * preco_un, 2)
                    diff = float(it.diferenca if it.diferenca is not None else (qtd_cnt - qtd_sis))
                    val_diff = round(diff * preco_un, 2)
                else:
                    qtd_cnt = None
                    val_cnt = None
                    diff = None
                    val_diff = None

                sit_str = it.situacao.value if hasattr(it.situacao, 'value') else str(it.situacao or "-")
                motivo_raw = it.motivo_ajuste.value if it.motivo_ajuste and hasattr(it.motivo_ajuste, 'value') else (str(it.motivo_ajuste) if it.motivo_ajuste else "-")
                contador_nome = str(it.contador.name) if it.contador else "-"
                contado_em_str = it.contado_em.strftime('%d/%m/%Y %H:%M') if it.contado_em else "-"
                obs_str = str(it.observacao or "-")

                # Valores da linha
                row_vals = [
                    (cod, "left"),
                    (nome, "left"),
                    (tipo_item, "center"),
                    (un_sigla, "center"),
                    (preco_un, "money"),
                    (qtd_sis, "num"),
                    (val_sis, "money"),
                    (qtd_cnt if qtd_cnt is not None else "-", "num" if qtd_cnt is not None else "center"),
                    (val_cnt if val_cnt is not None else "-", "money" if val_cnt is not None else "center"),
                    (diff if diff is not None else "-", "num" if diff is not None else "center"),
                    (val_diff if val_diff is not None else "-", "money" if val_diff is not None else "center"),
                    (sit_str, "center"),
                    (motivo_raw, "left"),
                    (contador_nome, "left"),
                    (contado_em_str, "center"),
                    (obs_str, "left")
                ]

                is_falta = (diff is not None and diff < 0)
                is_sobra = (diff is not None and diff > 0)

                for c_idx, (val, fmt) in enumerate(row_vals, 1):
                    cell = ws.cell(row=curr_row, column=c_idx)
                    cell.border = thin_border
                    
                    if is_falta and c_idx in [10, 11, 12]:
                        cell.fill = falta_fill
                        cell.font = falta_font
                    elif is_sobra and c_idx in [10, 11, 12]:
                        cell.fill = sobra_fill
                        cell.font = sobra_font
                    else:
                        cell.font = regular_font

                    if fmt == "money" and isinstance(val, (int, float)):
                        cell.value = val
                        cell.number_format = '#,##0.00 "STN"'
                        cell.alignment = Alignment(horizontal="right")
                    elif fmt == "num" and isinstance(val, (int, float)):
                        cell.value = val
                        cell.number_format = '#,##0.00'
                        cell.alignment = Alignment(horizontal="right")
                    else:
                        cell.value = str(val)
                        cell.alignment = Alignment(horizontal=fmt if fmt in ["left", "center", "right"] else "left")

                curr_row += 1

            # Linha de Totais da Tabela
            ws.cell(row=curr_row, column=1, value="TOTAL GERAL").font = bold_font
            ws.cell(row=curr_row, column=7, value=f"=SUM(G4:G{curr_row-1})").number_format = '#,##0.00 "STN"'
            ws.cell(row=curr_row, column=7).font = bold_font
            ws.cell(row=curr_row, column=9, value=f"=SUM(I4:I{curr_row-1})").number_format = '#,##0.00 "STN"'
            ws.cell(row=curr_row, column=9).font = bold_font
            ws.cell(row=curr_row, column=11, value=f"=SUM(K4:K{curr_row-1})").number_format = '#,##0.00 "STN"'
            ws.cell(row=curr_row, column=11).font = bold_font

            # Ajustar Largura das Colunas
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # ==========================================
        # FOLHA 2 — DIVERGÊNCIAS (FALTAS E SOBRAS)
        # ==========================================
        ws_div = wb.create_sheet(title="Divergências")
        itens_div = [it for it in inventario.items if it.situacao in [SituacaoDivergencia.FALTA, SituacaoDivergencia.SOBRA]]
        preencher_aba_itens(ws_div, itens_div, "RELATÓRIO EXCLUSIVO DE DIVERGÊNCIAS (FALTAS E SOBRAS)")

        # ==========================================
        # FOLHA 3 — CONTAGEM GERAL (TODOS OS ITENS)
        # ==========================================
        ws_geral = wb.create_sheet(title="Contagem Geral")
        preencher_aba_itens(ws_geral, inventario.items, "CONFERÊNCIA GERAL DE TODOS OS ARTIGOS")

        # ==========================================
        # FOLHA 4 — APENAS FALTAS (PREJUÍZO)
        # ==========================================
        ws_faltas = wb.create_sheet(title="Apenas Faltas")
        itens_faltas = [it for it in inventario.items if it.situacao == SituacaoDivergencia.FALTA]
        preencher_aba_itens(ws_faltas, itens_faltas, "AUDITORIA DE FALTAS DE STOCK (PERDAS)")

        # ==========================================
        # FOLHA 5 — APENAS SOBRAS (EXCESSO)
        # ==========================================
        ws_sobras = wb.create_sheet(title="Apenas Sobras")
        itens_sobras = [it for it in inventario.items if it.situacao == SituacaoDivergencia.SOBRA]
        preencher_aba_itens(ws_sobras, itens_sobras, "AUDITORIA DE SOBRAS DE STOCK (EXCESSOS)")

        # Ajustar largura das colunas na folha de Resumo
        for col in ws_resumo.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_resumo.column_dimensions[col_letter].width = max(max_len + 4, 15)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
